from math import ceil

import qtawesome as fa
from PyQt6.QtCore import Qt, QAbstractTableModel, QDate, QEvent, QModelIndex, QTimer
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit, QPushButton, QTableView,
    QHeaderView, QFrame, QAbstractItemView, QMessageBox, QDateEdit, QComboBox,
    QFileDialog, QApplication
)

from db.read.audit_trail import SEARCH_COLUMNS, get_date_range, get_records

ROW_LIMIT = 1000

COLUMNS = ["Timestamp", "Workstation", "Action", "Details", "IP Address",
           "User MAC", "Workstation MAC"]
DETAILS_COL = 3
MIN_DETAILS = 240
FIXED_WIDTHS = {0: 185, 1: 170, 2: 110, 4: 115, 5: 150, 6: 150}

# Excel export
XLSX_MIN_WIDTH = 10
XLSX_MAX_WIDTH = {DETAILS_COL: 70}     # every other column is capped at 40
XLSX_LINE_HEIGHT = 15.0                # points per line of 11 pt text

LEFT = Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter


def _icon(name, color, active=None):
    return fa.icon(name, color=color, color_active=active or color, color_disabled="#94A3B8")


def _qdate(d):
    return QDate(d.year, d.month, d.day)


def row_texts(db_row):
    """One database row -> list of display strings (also used by the Excel export)."""
    logged_at, workstation, action, details, ip, user_mac, ws_mac = db_row
    return [
        logged_at.strftime("%m/%d/%Y %I:%M:%S %p") if logged_at else "",
        workstation or "",
        action or "",
        details or "",
        ip or "",
        user_mac or "",
        ws_mac or "",
    ]


def write_xlsx(path, db_rows):
    """Bold, larger headers; column widths and row heights fitted to the content."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    table = [row_texts(r) for r in db_rows]
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Trail"

    header_font = Font(bold=True, size=13)
    header_align = Alignment(vertical="center", wrap_text=True)
    body_align = Alignment(vertical="top", wrap_text=True)

    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = header_font
        cell.alignment = header_align

    for r, texts in enumerate(table, start=2):
        for c, text in enumerate(texts, start=1):
            cell = ws.cell(row=r, column=c, value=text)
            cell.data_type = "s"            # text stays text, even if it starts with "="
            cell.alignment = body_align

    # Column widths: longest line in the column (or the header), within limits
    widths = []
    for c, name in enumerate(COLUMNS):
        longest = max((len(part) for texts in table for part in texts[c].split("\n")), default=0)
        width = max(longest, len(name) * 1.3) + 3
        width = min(max(width, XLSX_MIN_WIDTH), XLSX_MAX_WIDTH.get(c, 40))
        widths.append(width)
        ws.column_dimensions[get_column_letter(c + 1)].width = width

    # Row heights: number of wrapped lines in the tallest cell of the row
    ws.row_dimensions[1].height = 26
    for r, texts in enumerate(table, start=2):
        lines = 1
        for c, text in enumerate(texts):
            per_line = max(int(widths[c]) - 2, 1)
            n = sum(max(1, ceil(len(part) / per_line)) for part in text.split("\n"))
            lines = max(lines, n)
        ws.row_dimensions[r].height = min(lines * XLSX_LINE_HEIGHT + 4, 409)

    ws.freeze_panes = "A2"      # header stays visible while scrolling
    wb.save(path)


class AuditTableModel(QAbstractTableModel):
    """Read-only model. Display text is prepared once."""

    def __init__(self):
        super().__init__()
        self._rows = []

    def set_rows(self, db_rows):
        rows = [row_texts(r) for r in db_rows]
        self.beginResetModel()
        self._rows = rows
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self._rows)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(COLUMNS)

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if orientation == Qt.Orientation.Horizontal and role == Qt.ItemDataRole.DisplayRole:
            return COLUMNS[section]
        return None

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        if role == Qt.ItemDataRole.DisplayRole:
            return self._rows[index.row()][index.column()]
        if role == Qt.ItemDataRole.TextAlignmentRole:
            return LEFT
        return None


class AuditTrailWidget(QWidget):
    """Audit Trail: searchable, date-filtered log with Excel export."""

    def __init__(self):
        super().__init__()
        self.model = AuditTableModel()

        # Waits a moment after the user stops typing / picking before querying
        self._reload_timer = QTimer(self)
        self._reload_timer.setSingleShot(True)
        self._reload_timer.setInterval(350)
        self._reload_timer.timeout.connect(lambda: self.load_records(warn=False))

        # Re-measures row heights after the Details column changes width (word wrap)
        self._row_timer = QTimer(self)
        self._row_timer.setSingleShot(True)
        self._row_timer.setInterval(120)
        self._row_timer.timeout.connect(self.table_resize_rows)

        self._build_ui()

    # ── UI ────────────────────────────────────────────────────────────────
    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 24, 24, 24)
        root.setSpacing(16)

        # Heading
        title = QLabel("Audit Trail")
        title.setObjectName("HeadingTitle")
        subtitle = QLabel("Who did what, from which workstation, and when")
        subtitle.setObjectName("SubtitleLabel")
        head = QVBoxLayout()
        head.setSpacing(2)
        head.addWidget(title)
        head.addWidget(subtitle)
        root.addLayout(head)

        # Search & filter card
        card, body = self._card("SEARCH & FILTER")
        row = QHBoxLayout()
        row.setSpacing(16)

        self.cmb_column = QComboBox()
        for label, key in SEARCH_COLUMNS:
            self.cmb_column.addItem(label, key)

        self.txt_search = QLineEdit()
        self.txt_search.setPlaceholderText("Type to search...")
        self.txt_search.setClearButtonEnabled(True)

        # The real dates are set from the audit trail itself in _sync_dates()
        self.date_from = QDateEdit(QDate.currentDate())
        self.date_from.setCalendarPopup(True)
        self.date_from.setDisplayFormat("MM/dd/yyyy")

        self.date_to = QDateEdit(QDate.currentDate())
        self.date_to.setCalendarPopup(True)
        self.date_to.setDisplayFormat("MM/dd/yyyy")

        self._auto_from = None      # the date _sync_dates last put in each field,
        self._auto_to = None        # so a date picked by hand can be told apart

        row.addLayout(self._field("Search In", self.cmb_column), 1)
        row.addLayout(self._field("Search", self.txt_search), 3)
        row.addLayout(self._field("From", self.date_from), 1)
        row.addLayout(self._field("To", self.date_to), 1)
        body.addLayout(row)
        root.addWidget(card)

        self.txt_search.returnPressed.connect(lambda: self.load_records())
        self.txt_search.textChanged.connect(self._schedule_reload)
        self.cmb_column.currentIndexChanged.connect(self._on_column_changed)
        self.date_from.dateChanged.connect(self._schedule_reload)
        self.date_to.dateChanged.connect(self._schedule_reload)

        # Table
        self.table = QTableView()
        self.table.setModel(self.model)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setMinimumSectionSize(36)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setAlternatingRowColors(True)
        self.table.setShowGrid(False)
        self.table.setWordWrap(True)
        self.table.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.table.setVerticalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)

        hdr = self.table.horizontalHeader()
        hdr.setStretchLastSection(False)
        for col in range(len(COLUMNS)):
            hdr.setSectionResizeMode(col, QHeaderView.ResizeMode.Interactive)
        for col, width in FIXED_WIDTHS.items():
            self.table.setColumnWidth(col, width)
        self.table.setColumnWidth(DETAILS_COL, 400)

        # Details takes whatever width is left (never below MIN_DETAILS); wrapped
        # rows are re-measured whenever that width changes.
        self.table.viewport().installEventFilter(self)
        hdr.sectionResized.connect(self._on_section_resized)
        root.addWidget(self.table, 1)

        # Footer
        footer = QHBoxLayout()
        footer.setSpacing(8)
        self.lbl_count = QLabel("")
        self.lbl_count.setObjectName("SubtitleLabel")
        footer.addWidget(self.lbl_count)
        footer.addStretch(1)

        self.btn_export = QPushButton(" Export", icon=_icon("fa5s.download", "white"))
        self.btn_export.setObjectName("SecondaryButton")
        self.btn_export.clicked.connect(self.export_excel)

        btn_refresh = QPushButton(" Refresh", icon=_icon("fa5s.sync-alt", "white"))
        btn_refresh.setObjectName("PrimaryButton")
        btn_refresh.clicked.connect(self.refresh)

        footer.addWidget(self.btn_export)
        footer.addWidget(btn_refresh)
        root.addLayout(footer)

    def _card(self, title):
        frame = QFrame()
        frame.setObjectName("Card")
        body = QVBoxLayout(frame)
        body.setContentsMargins(8, 8, 8, 8)
        body.setSpacing(12)
        lbl = QLabel(title)
        lbl.setObjectName("CardTitle")
        body.addWidget(lbl)
        return frame, body

    def _field(self, caption, widget):
        box = QVBoxLayout()
        box.setSpacing(4)
        lbl = QLabel(caption)
        lbl.setObjectName("FieldLabel")
        box.addWidget(lbl)
        box.addWidget(widget)
        return box

    # ── Column sizing / word wrap ─────────────────────────────────────────
    def eventFilter(self, obj, event):
        if obj is self.table.viewport() and event.type() == QEvent.Type.Resize:
            self._fit_details()
        return super().eventFilter(obj, event)

    def _fit_details(self):
        hdr = self.table.horizontalHeader()
        others = sum(hdr.sectionSize(c) for c in range(len(COLUMNS)) if c != DETAILS_COL)
        free = self.table.viewport().width() - others
        self.table.setColumnWidth(DETAILS_COL, max(MIN_DETAILS, free))

    def _on_section_resized(self, index, _old, _new):
        if index == DETAILS_COL:
            self._row_timer.start()

    def table_resize_rows(self):
        self.table.resizeRowsToContents()

    # ── Dates ─────────────────────────────────────────────────────────────
    def _sync_dates(self):
        """From / To follow the first and last day in the audit trail, and the
        calendars only allow days inside that range. A date the user picked by
        hand is kept (as long as it is still inside the range)."""
        try:
            first, last = get_date_range()
        except Exception:
            return      # load_records will show the database error

        today = QDate.currentDate()
        low = _qdate(first) if first else today
        high = _qdate(last) if last else today

        for edit, attr, value in ((self.date_from, "_auto_from", low),
                                  (self.date_to, "_auto_to", high)):
            auto = getattr(self, attr)
            follow = auto is None or edit.date() == auto    # not changed by the user
            edit.blockSignals(True)
            edit.setDateRange(low, high)
            if follow:
                edit.setDate(value)
            edit.blockSignals(False)
            setattr(self, attr, value)

    def refresh(self):
        self._sync_dates()
        self.load_records()

    # ── Data ──────────────────────────────────────────────────────────────
    def showEvent(self, event):
        super().showEvent(event)
        if not event.spontaneous():     # shown by navigation, not by minimize/restore
            self._sync_dates()
            self.load_records(warn=False)

    def _schedule_reload(self, *_):
        self._reload_timer.start()

    def _on_column_changed(self, *_):
        if self.txt_search.text().strip():
            self._schedule_reload()

    def _read_filters(self, warn):
        d_from = self.date_from.date().toPyDate()
        d_to = self.date_to.date().toPyDate()
        if d_from > d_to:
            if warn:
                QMessageBox.warning(self, "Invalid Dates",
                                    "The 'From' date is later than the 'To' date.")
            else:
                self.lbl_count.setText("The 'From' date is later than the 'To' date.")
            return None
        return {
            "search": self.txt_search.text().strip(),
            "column": self.cmb_column.currentData(),
            "date_from": d_from,
            "date_to": d_to,
        }

    def load_records(self, warn=True):
        self._reload_timer.stop()
        filters = self._read_filters(warn)
        if filters is None:
            return

        try:
            rows = get_records(limit=ROW_LIMIT, **filters)
        except Exception as e:
            QMessageBox.critical(self, "Database Error", f"Could not load the audit trail:\n{e}")
            return

        self.table.clearSelection()
        self.model.set_rows(rows)
        self.table.resizeRowsToContents()

        suffix = f" (showing newest {ROW_LIMIT:,}, export for all)" if len(rows) >= ROW_LIMIT else ""
        self.lbl_count.setText(f"{len(rows):,} record(s){suffix}")
        self.btn_export.setEnabled(len(rows) > 0)

    # ── Export ────────────────────────────────────────────────────────────
    def export_excel(self):
        """Exports every record matching the current filters (not just the first 1,000)."""
        filters = self._read_filters(warn=True)
        if filters is None:
            return

        try:
            import openpyxl  # noqa: F401
        except ImportError:
            QMessageBox.warning(self, "Missing Package",
                                "Excel export needs the 'openpyxl' package.\n\n"
                                "Install it with:  pip install openpyxl")
            return

        default = f"audit_trail_{filters['date_from']:%Y%m%d}_{filters['date_to']:%Y%m%d}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "Export Audit Trail", default,
                                              "Excel Files (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            rows = get_records(limit=None, **filters)
            write_xlsx(path, rows)
        except Exception as e:
            QApplication.restoreOverrideCursor()
            QMessageBox.critical(self, "Export Failed",
                                 f"Could not export the audit trail:\n{e}\n\n"
                                 "If the file is open in Excel, close it and try again.")
            return
        QApplication.restoreOverrideCursor()

        QMessageBox.information(self, "Exported", f"{len(rows):,} record(s) saved to:\n{path}")